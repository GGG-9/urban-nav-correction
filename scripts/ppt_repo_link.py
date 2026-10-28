"""Put the GitHub repository link on the closing slide of the group deck.

    python scripts/ppt_repo_link.py

Reads/writes `E:\\大数据ppt\\第三组.pptx` (a timestamped backup is written
next to it first), adds a label + URL + QR code under the "Thank You" block
on the last slide and appends one line to that slide's speaker notes.

Fonts and colours are taken from the deck itself: 微软雅黑 for Latin,
34765E for the small tracked labels, 31481D for the dark green.
"""
from __future__ import annotations

import datetime as dt
import pathlib
import shutil

import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

DECK = pathlib.Path(r"E:\大数据ppt\第三组.pptx")
URL = "https://github.com/GGG-9/urban-nav-correction"
DISPLAY = "github.com/GGG-9/urban-nav-correction"
QR_PNG = pathlib.Path(r"C:\Users\m1775\AppData\Local\Temp\repo_qr.png")

INK = RGBColor(0x31, 0x48, 0x1D)      # dark green, deck accent2
TEAL = RGBColor(0x34, 0x76, 0x5E)     # deck accent4, small labels
MUTE = RGBColor(0x66, 0x75, 0x6B)     # grey-green for the caption
FONT = "微软雅黑"


def tracked(run, spc: int) -> None:
    """Letter spacing, in 1/100 pt, as the deck's labels use."""
    run.font._rPr.set("spc", str(spc))


def textbox(slide, x, y, w, h, text, size, color, bold=False, spc=None,
            align=None, link=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    f = r.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    if spc is not None:
        tracked(r, spc)
    if link:
        r.hyperlink.address = link
    return box


def make_qr(path: pathlib.Path) -> None:
    qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M,
                       box_size=14, border=2)
    qr.add_data(URL)
    qr.make(fit=True)
    qr.make_image(fill_color="#31481D", back_color="white").save(path)


def main() -> None:
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = DECK.with_name(f"{DECK.stem}_备份_{stamp}{DECK.suffix}")
    shutil.copy2(DECK, backup)
    print("backup:", backup.name)

    make_qr(QR_PNG)

    prs = Presentation(DECK)
    s = prs.slides[-1]                      # Questions & Discussion / Thank You

    textbox(s, 1.16, 5.24, 4.20, 0.24, "CODE · DATA · FIGURES", 10, TEAL,
            bold=True, spc=150)
    # no hyperlink: PowerPoint repaints linked runs in the theme's hyperlink
    # green, which clashes with the deck palette; the QR carries the "open"
    textbox(s, 1.13, 5.44, 8.10, 0.55, DISPLAY, 22, INK, bold=True)
    textbox(s, 1.16, 6.04, 7.40, 0.28,
            "Source, simulator, configs and every figure in this deck.",
            11, MUTE)

    s.shapes.add_picture(str(QR_PNG), Inches(9.98), Inches(5.06),
                         Inches(1.70), Inches(1.70))
    textbox(s, 9.98, 6.80, 1.70, 0.20, "SCAN TO OPEN", 8.5, MUTE, spc=100,
            align=PP_ALIGN.CENTER)

    notes = s.notes_slide.notes_text_frame
    notes.text = notes.text.rstrip() + (
        "\nThe full source, the 48-route simulator and every figure in this "
        "deck are on GitHub — the QR code on this slide opens the repository."
    )

    prs.save(DECK)
    print("saved:", DECK)


if __name__ == "__main__":
    main()
