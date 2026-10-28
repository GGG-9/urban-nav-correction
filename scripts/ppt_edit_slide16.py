"""Rewrite slide 17 (index 16, "3.4 TRAINING & SAFETY GATING") of the deck:
the three text cards become three figure cards, and the numbers are the ones
`scripts/ppt_figs.py` measures.

    python scripts/ppt_edit_slide16.py [deck.pptx]

The original is copied to ``<deck>.bak.pptx`` before the file is overwritten;
no other slide is touched.
"""
from __future__ import annotations

import copy
import pathlib
import shutil
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = pathlib.Path(__file__).resolve().parents[1]
FIG = ROOT / "results" / "ppt"
DECK = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else pathlib.Path(
    r"E:\大作业\GPU-Parallel Batched Kalman Smoothing and LoRA Fine-Tuned "
    r"Time-Series Foundation Model.pptx")

SLIDE = 16
CARD_L = (0.62, 4.74, 8.85)
TEXT_L = (0.88, 4.99, 9.10)

LABELS = ["STRATEGY 01 · IMBALANCE LOSS",
          "STRATEGY 02 · TRAINING OUTCOME",
          "STRATEGY 03 · SAFETY FALLBACK"]
HEADLINES = ["×22 on the 7 % that matters",
             "Converged — Never Beats Zero",
             "Validation Pays, Test Punishes"]
CAPTIONS = [
    [("Windows ×2, steps ×11 → ", False), ("×22", True),
     (" per outage sample. 7.0 % of samples, 17.1 % of the residual, ", False),
     ("81.7 %", True), (" of the gradient.", False)],
    [("Train loss 5.30 → 4.40 (LSTM), 4.98 → 1.94 (GPT4TS). Best held-out ", False),
     ("11.79 / 11.68 m²", True),
     (", never under the ", False), ("11.48", True),
     (" of predicting zero.", False)],
    [("Unguarded α* = 2.00 / 1.90 buys +2.7 % on validation and ", False),
     ("−2.2 %", True),
     (" on test; the one-SE rule keeps ", False), ("α = 0.00", True),
     (".", False)],
]
FOOTNOTE = ("Mixed precision kept as an implementation detail: FP16 autocast forward, "
            "loss and gradient clipping in FP32 (unscale_ before clipping), "
            "AdamW lr 1e-4 + OneCycleLR.")


def shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def place(sh, left=None, top=None, width=None, height=None):
    if left is not None:
        sh.left = Inches(left)
    if top is not None:
        sh.top = Inches(top)
    if width is not None:
        sh.width = Inches(width)
    if height is not None:
        sh.height = Inches(height)


def set_rich(tf, parts, size=11.5):
    """Replace every run in the text frame with ``parts`` = [(text, bold), ...],
    cloning the formatting of the first run so font and colour survive."""
    p = tf.paragraphs[0]
    body = tf._txBody
    for extra in body.findall(qn("a:p"))[1:]:
        body.remove(extra)
    runs = list(p.runs)
    if not runs:
        return
    template = copy.deepcopy(runs[0]._r)
    for r in runs:
        r._r.getparent().remove(r._r)
    for text, bold in parts:
        el = copy.deepcopy(template)
        el.find(qn("a:t")).text = text
        rpr = el.find(qn("a:rPr"))
        if rpr is None:
            rpr = el.makeelement(qn("a:rPr"), {})
            el.insert(0, rpr)
        rpr.set("b", "1" if bold else "0")
        rpr.set("sz", str(int(size * 100)))
        p._p.append(el)


def set_text(sh, text, size=None, bold=None):
    p = sh.text_frame.paragraphs[0]
    runs = list(p.runs)
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    r = runs[0]
    r.text = text
    if size is not None:
        r.font.size = Pt(size)
    if bold is not None:
        r.font.bold = bold


def main() -> None:
    prs = Presentation(DECK)
    slide = prs.slides[SLIDE]
    print(f"{DECK.name}  slide {SLIDE}: " +
          shape(slide, "TextBox 6").text_frame.text)

    # 1. cards grow upward into the empty band under the title
    for i, nm in enumerate(("Rounded Rectangle 2", "Rounded Rectangle 3",
                            "Rounded Rectangle 4")):
        place(shape(slide, nm), top=1.32, height=4.18)

    # 2. labels / headlines / captions
    for i, (lab, head, cap) in enumerate(zip(LABELS, HEADLINES, CAPTIONS)):
        L = TEXT_L[i]
        lb = shape(slide, ("TextBox 9", "TextBox 12", "TextBox 15")[i])
        place(lb, top=1.55)
        set_text(lb, lab)
        hd = shape(slide, ("TextBox 10", "TextBox 13", "TextBox 16")[i])
        place(hd, top=1.79, height=0.32)
        set_text(hd, head)
        cp = shape(slide, ("TextBox 11", "TextBox 14", "TextBox 17")[i])
        place(cp, left=L, top=4.58, width=3.47, height=0.74)
        set_rich(cp.text_frame, cap, size=11.5)

    # 3. the three figures
    for i, png in enumerate(("a_imbalance.png", "b_training.png", "c_gate.png")):
        slide.shapes.add_picture(str(FIG / png), Inches(TEXT_L[i]),
                                 Inches(2.24), width=Inches(3.47),
                                 height=Inches(2.24))

    # 4. formula bar: corrected equation, shorter lead-in
    bar = shape(slide, "Rounded Rectangle 1")
    place(bar, top=5.62, height=0.62)
    lbl = shape(slide, "TextBox 8")
    place(lbl, top=5.81, width=7.0, height=0.31)
    set_text(lbl, "Validation shrinkage safety gate:")
    old = shape(slide, "Picture 5")
    old._element.getparent().remove(old._element)
    slide.shapes.add_picture(str(FIG / "formula.png"), Inches(8.48), Inches(5.78),
                             width=Inches(3.94), height=Inches(0.30))

    # 5. mixed precision -> one-line footnote under the bar
    tb = slide.shapes.add_textbox(Inches(0.62), Inches(6.36), Inches(12.08),
                                  Inches(0.30))
    tf = tb.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = FOOTNOTE
    r.font.size = Pt(10)
    r.font.name = "DM Sans"
    from pptx.dml.color import RGBColor
    r.font.color.rgb = RGBColor(0x71, 0x71, 0x7A)

    backup = DECK.with_name(DECK.stem + ".bak.pptx")
    shutil.copy2(DECK, backup)
    try:
        prs.save(DECK)
        out = DECK
    except PermissionError:                     # open in PowerPoint: write beside it
        out = DECK.with_name(DECK.stem + "-v2.pptx")
        prs.save(out)
    print(f"backup -> {backup.name}")
    print(f"saved  -> {out.name}")


if __name__ == "__main__":
    main()
