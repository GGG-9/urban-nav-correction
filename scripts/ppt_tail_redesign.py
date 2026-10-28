"""Rebuild slides 18-23 of the deck from scratch (3.5 through the closing page).

    python scripts/ppt_tail_redesign.py [deck.pptx]

Input is the v3 deck (slide 17 already redesigned). The old tail -- 3.5, the
chapter-04 divider, the map-prior/summary page, the thank-you and the image
credits -- was written as prose cards with several numbers that this project
later measured differently. Every number below comes from results/summary.md,
results/bench.json or the run jsons; the pages are laid out as argument left /
evidence right, matching slide 17.

New order (1-based):
    18  3.5  the ladder + the information wall
    19  CH04 divider
    20  4.1  road-centreline projection        (fig7)
    21  4.2  lane-level mapped result          (fig8)
    22  4.3  three numbers, three limits
    23  Thank You (+ image credits, folded in from the old last page)
"""
from __future__ import annotations

import copy
import pathlib
import shutil
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = pathlib.Path(__file__).resolve().parents[1]
FIG = ROOT / "results" / "ppt"
RES = ROOT / "results"
DECK = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else pathlib.Path(
    r"E:\大作业\GPU-Parallel Batched Kalman Smoothing and LoRA Fine-Tuned "
    r"Time-Series Foundation Model-v3.pptx")

FIRST = 17                     # 0-based index of the first rebuilt page (slide 18)
N_OLD = 5                      # old slides 18..22 -> indices 17..21
N_NEW = 6

FONT = "DM Sans"
INK = RGBColor(0x18, 0x18, 0x1B)
MUTE = RGBColor(0x57, 0x53, 0x4E)
FAINT = RGBColor(0x71, 0x71, 0x7A)
ORANGE = RGBColor(0xEA, 0x58, 0x0C)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
GREY = RGBColor(0xA8, 0xA2, 0x9E)
HAIR = RGBColor(0xE7, 0xE5, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DIM = RGBColor(0xA1, 0xA1, 0xAA)
DARK = RGBColor(0x18, 0x18, 0x1B)


# --------------------------------------------------------------------------- #
# primitives
# --------------------------------------------------------------------------- #
def text(S, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line=None, space_after=0):
    """``lines`` = one paragraph per list entry; a paragraph is [(text, kw), ...]."""
    if lines and isinstance(lines[0], tuple):
        lines = [lines]
    tb = S.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, parts in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line is not None:
            p.line_spacing = line
        if space_after:
            p.space_after = Pt(space_after)
        for s, kw in parts:
            r = p.add_run()
            r.text = s
            f = r.font
            f.name = FONT
            f.size = Pt(kw.get("size", 11))
            f.bold = kw.get("bold", False)
            f.color.rgb = kw.get("color", MUTE)
            if kw.get("spc"):
                r._r.get_or_add_rPr().set("spc", str(kw["spc"]))
    return tb


def rect(S, x, y, w, h, color, rounded=False, radius=0.5):
    shp = S.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def pic(S, name, x, y, w, h):
    return S.shapes.add_picture(str(FIG / name), Inches(x), Inches(y),
                                width=Inches(w), height=Inches(h))


def header(S, eyebrow, title, color):
    text(S, 0.62, 0.46, 12.71, 0.19, [(eyebrow, {"size": 10.5, "bold": True, "color": color})])
    text(S, 0.62, 0.71, 12.71, 0.44, [(title, {"size": 24, "bold": True, "color": INK})])


def sec(S, x, y, num, label, color, w=6.0):
    text(S, x, y, w, 0.20,
         [(num, {"size": 10.5, "bold": True, "color": GREY, "spc": 120}),
          ("   " + label, {"size": 10.5, "bold": True, "color": color, "spc": 120})])


def strip(S, y, head, tail, x=0.62, w=12.08):
    rect(S, x, y, w, 0.44, DARK, rounded=True, radius=0.5)
    text(S, x + 0.30, y + 0.08, w - 0.6, 0.30,
         [(head, {"size": 10.5, "bold": True, "color": WHITE, "spc": 120}),
          ("    " + tail, {"size": 10, "color": DIM})],
         anchor=MSO_ANCHOR.MIDDLE)


def stat_row(S, x, y, w, big, body, big_w=1.35, color=INK):
    # numeral runs stay un-bolded on purpose: DM Sans Bold swaps in a tailed
    # zero whenever "0" precedes a period ("0.9" -> looks like "Q.9")
    text(S, x, y, big_w, 0.55, [(big, {"size": 30, "color": color})],
         anchor=MSO_ANCHOR.MIDDLE)
    text(S, x + big_w + 0.18, y + 0.02, w - big_w - 0.18, 0.70,
         [(body, {"size": 10.5, "color": MUTE})], line=1.18)


# --------------------------------------------------------------------------- #
# 18 -- 3.5  the ladder + the wall
# --------------------------------------------------------------------------- #
def page_35(S):
    header(S, "CHAPTER 03: FOUNDATION MODEL & LORA · 3.5 WHAT IT ACTUALLY BOUGHT",
           "What Did the Learned Corrector Buy — and Where Does It Stop?", ORANGE)

    sec(S, 0.62, 1.40, "01", "THE LADDER", ORANGE)
    pic(S, "ladder.png", 0.62, 1.66, 5.43, 3.10)
    text(S, 0.62, 4.88, 5.43, 1.05,
         [("B1 → B3 removes 90.5 % of the outage error; the worst single outage falls "
           "413 → 28.3 m. Running the same smoother on GPU windows costs +8.2 % "
           "(10.75 m) — the price of seeing ±90 s instead of the whole route.",
           {"size": 10.5, "color": MUTE})], line=1.2)

    sec(S, 6.35, 1.40, "02", "THE WALL", ORANGE)
    rows = [
        ("84 %", "of the RTS residual is a smooth quadratic in outage time — a frozen "
                 "initial-state error, not noise."),
        ("91 %", "of the outage error would vanish if that smooth part were known "
                 "exactly, i.e. if the missing fix existed."),
        ("0.9 %", "is all an oracle IMU calibration would add. The bias is not what is "
                  "left over."),
        ("0.0 %", "is what the trained corrector recovers on unseen routes: the channels "
                  "do not carry the missing state."),
    ]
    for i, (big, body) in enumerate(rows):
        stat_row(S, 6.35, 1.68 + 0.88 * i, 6.35, big, body)

    text(S, 6.35, 5.32, 6.30, 0.90,
         [("Direction transfers across routes (r ≈ 0.70); amplitude does not. The size "
           "of the offset is fixed at outage entry, where no fix is recorded — a larger "
           "network cannot add that information.", {"size": 10.5, "color": FAINT})],
         line=1.2)

    strip(S, 6.68, "THE GATE IS THE RESULT",
          "α = 0 on 8 / 8 test routes — the module may only act where it demonstrably helps")


# --------------------------------------------------------------------------- #
# 19 -- chapter 04 divider
# --------------------------------------------------------------------------- #
def page_divider(S):
    rect(S, 0.83, 5.17, 0.83, 0.06, TEAL)
    text(S, 0.83, 2.27, 12.50, 0.27, [("CHAPTER 04", {"size": 15, "bold": True, "color": TEAL})])
    text(S, 0.83, 2.67, 12.36, 1.31,
         [("Spatial Map Prior & Lane-Level Mapping", {"size": 36, "bold": True, "color": INK})])
    text(S, 0.83, 4.23, 11.56, 0.60,
         [("4.1 Road-Centreline Projection  ➔  4.2 Lane-Level Mapped Result  ➔  "
           "4.3 Three Numbers, Three Limits", {"size": 16.5, "color": RGBColor(0x3F, 0x3F, 0x46)})])


# --------------------------------------------------------------------------- #
# 20 -- 4.1 map prior
# --------------------------------------------------------------------------- #
def page_map(S):
    header(S, "CHAPTER 04: MAP PRIOR & LANE-LEVEL MAPPING · 4.1 ROAD-CENTRELINE PROJECTION",
           "Can a Road Map Remove What the Inertial Filter Cannot See?", TEAL)

    hero = [(0.62, 4.10, "9.93 → 6.30 m", "test-mean outage error, B3 → map prior (−36.5 %)"),
            (4.95, 3.45, "8 / 8", "test routes improve, +26 % to +46 %"),
            (8.55, 4.15, "±1 m", "survey error tolerated; 2 m breaks route 46")]
    for x, w, big, small in hero:
        text(S, x, 1.30, w, 0.46, [(big, {"size": 26, "color": INK})])
        text(S, x, 1.76, w, 0.26, [(small, {"size": 9, "color": MUTE})])

    pic(S, "../fig7_map_prior_b3.png", 0.62, 2.14, 12.08, 3.30)
    text(S, 0.62, 5.56, 12.08, 0.62,
         [("Shown on the numpy reference (B3). The centreline supplies the cross-track "
           "component no inertial sensor can observe and leaves the along-track part alone; "
           "the GPU trajectory follows the same path, 10.75 → 6.98 m.",
           {"size": 10.5, "color": MUTE})], line=1.2)
    text(S, 0.62, 6.30, 12.08, 0.55,
         [("Projection is continuity-constrained — arc length monotone, speed limit, "
           "heading term (scripts/map_prior.py) — not nearest-point: nearest-point reaches "
           "only 7.52 m and regresses the self-intersecting route 46 by 36 %.",
           {"size": 9.5, "color": FAINT})], line=1.18)


# --------------------------------------------------------------------------- #
# 21 -- 4.2 lane-level result
# --------------------------------------------------------------------------- #
def page_lane(S):
    header(S, "CHAPTER 04: MAP PRIOR & LANE-LEVEL MAPPING · 4.2 WHAT IT MEANS FOR MAPPING",
           "Does the Corrected Trajectory Land in the Right Lane?", TEAL)

    text(S, 0.62, 1.30, 7.60, 0.46,
         [("43.9 → 6.5 → 0.28 m", {"size": 26, "color": INK})])
    text(S, 0.62, 1.76, 7.60, 0.26,
         [("lane-line lateral offset: dead reckoning → filter → map prior",
           {"size": 9, "color": MUTE})])
    text(S, 8.85, 1.30, 3.85, 0.46, [("86 %", {"size": 26, "color": INK})])
    text(S, 8.85, 1.76, 3.85, 0.26,
         [("of mapped points inside the 0.5 m HD tolerance", {"size": 9, "color": MUTE})])

    pic(S, "../fig8_laneline.png", 0.62, 2.14, 12.08, 3.37)
    text(S, 0.62, 5.62, 12.08, 0.62,
         [("A lane line built from these poses lands on the road only after the map prior: "
           "43.9 m of lateral offset from dead reckoning, 6.5 m from the filter, 0.28 m with "
           "the prior — and 4.1 % / 6.5 % / 86.2 % of points inside the 0.5 m tolerance.",
           {"size": 10.5, "color": MUTE})], line=1.2)
    text(S, 0.62, 6.36, 12.08, 0.50,
         [("0.28 m is the chord error of a 1 Hz centreline, not a pose accuracy — with a "
           "prior map this measures map update, not mapping from scratch. The along-track "
           "component (40.7 → 6.9 m) is untouched by design.", {"size": 9.5, "color": FAINT})],
         line=1.18)


# --------------------------------------------------------------------------- #
# 22 -- 4.3 conclusion
# --------------------------------------------------------------------------- #
def page_conclusion(S):
    header(S, "CHAPTER 04: MAP PRIOR & LANE-LEVEL MAPPING · 4.3 THREE NUMBERS, THREE LIMITS",
           "What Do These Three Numbers Add Up To?", TEAL)

    cols = [
        ("48.6×", "THROUGHPUT",
         "1 M IMU rows: 2,245 s on CPU vs 46.2 s on one RTX 4090; 177.6× at chunk 500. "
         "The full 8.64 M-row corpus filters in 17 min."),
        ("90.5 %", "ACCURACY",
         "Outage error 104.65 → 9.93 m and the worst outage 413 → 28.3 m. The windowed "
         "GPU pipeline pays 8.2 % of that for its parallelism."),
        ("0.28 m", "MAPPING",
         "Lane-line lateral offset after the map prior; 86 % of mapped points fall inside "
         "the 0.5 m HD tolerance."),
    ]
    for i, (big, lab, body) in enumerate(cols):
        x = 0.62 + 4.18 * i
        text(S, x, 1.32, 3.73, 0.80, [(big, {"size": 46, "color": INK})])
        text(S, x, 2.16, 3.73, 0.22,
             [(lab, {"size": 9, "bold": True, "color": TEAL, "spc": 140})])
        text(S, x, 2.44, 3.73, 1.55, [(body, {"size": 10.5, "color": MUTE})], line=1.22)

    rect(S, 0.62, 4.22, 12.08, 0.012, HAIR)
    text(S, 0.62, 4.36, 12.08, 0.22,
         [("WHAT WE WOULD TEST NEXT", {"size": 9, "bold": True, "color": GREY, "spc": 140})])

    bullets = [
        ("Replay a real IMU / GNSS log.", " The corpus here is simulated; the filter and "
         "the prior do not change, only the noise model does."),
        ("Widen the window.", " 90 s of burn-in is a requirement, not the window length — "
         "a longer window recovers most of the 8.2 %."),
        ("Feed the corrector the missing fix.", " Until a sensor supplies what the outage "
         "removed, the gate keeps α at zero and the module stays inert."),
    ]
    for i, (head, rest) in enumerate(bullets):
        y = 4.70 + 0.42 * i
        text(S, 0.62, y, 0.20, 0.26, [("•", {"size": 11, "color": TEAL})])
        text(S, 0.90, y, 11.80, 0.34,
             [(head, {"size": 11, "bold": True, "color": INK}),
              (rest, {"size": 11, "color": MUTE})])

    strip(S, 6.62, "REPRODUCIBLE",
          "48 routes · 8.64 M rows · every figure re-renders from results/summary.md")


# --------------------------------------------------------------------------- #
# 23 -- thank you
# --------------------------------------------------------------------------- #
def page_thanks(S):
    text(S, 0.0, 2.62, 13.333, 0.25,
         [("SOUTHEAST UNIVERSITY · ENGINEERING BIG-DATA ANALYSIS",
           {"size": 12, "bold": True, "color": ORANGE})], align=PP_ALIGN.CENTER)
    text(S, 0.0, 3.06, 13.333, 0.90,
         [("Thank You · Questions & Discussion", {"size": 42, "bold": True, "color": RGBColor(0xFA, 0xFA, 0xFA)})],
         align=PP_ALIGN.CENTER)
    text(S, 0.0, 4.06, 13.333, 0.34,
         [("GPU-Parallel Batched Kalman Smoothing & LoRA Fine-Tuned Time-Series Foundation Model",
           {"size": 15, "color": DIM})], align=PP_ALIGN.CENTER)
    text(S, 0.0, 6.86, 13.333, 0.24,
         [("Image credits: oxts.com (GNSS/INS test vehicle) · pexels.com (server hardware)",
           {"size": 9, "color": RGBColor(0x52, 0x52, 0x5B)})], align=PP_ALIGN.CENTER)


BUILDERS = [page_35, page_divider, page_map, page_lane, page_conclusion, page_thanks]

# Slides 1-17 are copied through untouched -- except where an AI-written number
# contradicts the measurements that were redone afterwards (92 % vs the measured
# 90.5 %, and a 11.9 h / 167x claim that no run supports).
STALE = {
    "Fuses 100Hz IMU with 1Hz GNSS to cut 92% of outage drift (104.65m down to 9.93m).":
        "Fuses 100Hz IMU with 1Hz GNSS to cut 90.5% of outage drift (104.65m down to 9.93m).",
    "Batches 192 windows across 48 routes on CUDA to compress 11.9h serial runtime into 17min (up to 167x).":
        "Batches 192 windows across 48 routes on CUDA: the full 8.64M-row corpus "
        "filters in 17 min, 48.6x at 1M rows.",
    # "over 90 %" is literally true of every B1 -> B3 number in the deck (90.5 %
    # mean, 93.1 % worst outage) and avoids a bold "0." in display type, which
    # this font renders with a Q-like tailed zero
    "RTS vs. Forward KF: Why Full-State Smoothing Cuts Drift by 92%":
        "RTS vs. Forward KF: Why Full-State Smoothing Cuts Drift by Over 90%",
}


def fix_stale(prs) -> None:
    n = 0
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.text in STALE:
                        r.text = STALE[r.text]
                        n += 1
    print(f"stale numbers corrected: {n}")


# --------------------------------------------------------------------------- #
def main() -> None:
    prs = Presentation(DECK)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    old_els = ids[FIRST:FIRST + N_OLD]
    old_rids = [e.get(qn("r:id")) for e in old_els]
    old_bg = [prs.part.related_part(r)._element.find(qn("p:cSld")).find(qn("p:bg"))
              for r in old_rids]
    print(f"deck has {len(ids)} slides; replacing {N_OLD}, adding {N_NEW}")

    blank = next(lay for lay in prs.slide_layouts if lay.name == "Blank")
    made = []
    for i, build in enumerate(BUILDERS):
        S = prs.slides.add_slide(blank)
        S.part.partname = PackURI(f"/ppt/slides/slide{90 + i}.xml")   # collision-proof
        assert len(S.shapes) == 0
        made.append(S)

    # backgrounds: content pages keep the old 18 tint, the divider takes the old
    # divider's chapter tint (F0FDFA), the closing page the old dark one
    pick = [0, 1, 0, 0, 0, 3]
    for i, S in enumerate(made):
        bg = old_bg[pick[i]]
        if bg is not None:
            S._element.find(qn("p:cSld")).insert(0, copy.deepcopy(bg))

    for build, S in zip(BUILDERS, made):
        build(S)

    # place them in order, then drop the old pages
    for k, S in enumerate(made):
        el = [e for e in lst if prs.part.related_part(e.get(qn("r:id"))).partname
              == S.part.partname][0]
        lst.remove(el)
        lst.insert(FIRST + k, el)
    for e, r in zip(old_els, old_rids):
        lst.remove(e)
        prs.part.drop_rel(r)

    fix_stale(prs)

    print(f"slides now: {len(lst)}")
    backup = DECK.with_name(DECK.stem.replace("-v3", "") + "-v3.bak.pptx")
    shutil.copy2(DECK, backup)
    out = DECK.with_name(DECK.stem.replace("-v3", "") + "-v4.pptx")
    prs.save(out)
    print(f"backup -> {backup.name}\nsaved  -> {out.name}")


if __name__ == "__main__":
    main()
