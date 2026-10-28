"""The speaker's own part (deck pages 17-23) rebuilt on the green course template.

    python scripts/ppt_green_deck.py

Reads   E:\\大作业\\课程汇报绿色模板.pptx       (5 sample slides, two masters)
Writes  E:\\大作业\\课程汇报绿色模板_我的部分.pptx  (7 pages, English)

Only the second master carries the green design: its "节标题" layout draws the
green band, the SEU seal and the page-number tab, and its "标题和内容" layout
draws the dark-green left panel used by the chapter divider.  Palette is taken
from that master's theme (name "自定义 6"):

    accent1 587558   accent2 31481D   accent3 9DC48C   accent4 34765E
    accent5 BACF85   background lt2 F6F8F9

The figures in results/ppt/green/ were rendered by the same scripts as the main
deck with FIG_PALETTE=green, so every number on the slides is the one that is
checked in results/summary.md.
"""
from __future__ import annotations

import pathlib
import shutil

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

TEMPLATE = pathlib.Path(r"E:\大作业\课程汇报绿色模板.pptx")
OUT = pathlib.Path(r"E:\大作业\课程汇报绿色模板_我的部分.pptx")
ASSETS = pathlib.Path(r"E:\大作业\urban-nav-correction\results\ppt\green")

# template palette
G_DARK = RGBColor(0x31, 0x48, 0x1D)
G_MAIN = RGBColor(0x58, 0x75, 0x58)
G_LIGHT = RGBColor(0x9D, 0xC4, 0x8C)
G_PALE = RGBColor(0xBA, 0xCF, 0x85)
G_TEAL = RGBColor(0x34, 0x76, 0x5E)
INK = RGBColor(0x1F, 0x23, 0x28)
MUTE = RGBColor(0x5D, 0x6B, 0x5D)
LINE = RGBColor(0xD5, 0xDC, 0xD5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WASH = RGBColor(0xEE, 0xF3, 0xEA)
BAND_TXT = RGBColor(0xE8, 0xEF, 0xE4)      # chapter tag on the green band

F = "微软雅黑"
F_LIGHT = "微软雅黑 Light"
F_SERIF = "Times New Roman"


# ----------------------------------------------------------------- helpers
def _style_run(run, size, color, bold, font, italic, spc):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = color
    if spc:
        run._r.get_or_add_rPr().set("spc", str(int(spc)))


def text(slide, x, y, w, h, paras, *, size=12, color=INK, bold=False, font=F,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0,
         space_after=0, italic=False, spc=None):
    """A plain text box.  ``paras`` is a string, a list of strings, or a list
    of runs ``[(txt, {"size": .., "color": ..}), ...]``; a paragraph may also
    be ``{"runs": [...], "align": ..}``."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        p.alignment = align
        if isinstance(spec, str):
            runs = [(spec, {})]
        elif isinstance(spec, dict):
            runs = spec["runs"]
            p.alignment = spec.get("align", align)
        else:
            runs = spec
        for t, st in runs:
            r = p.add_run()
            r.text = t
            _style_run(r, st.get("size", size), st.get("color", color),
                       st.get("bold", bold), st.get("font", font),
                       st.get("italic", italic), st.get("spc", spc))
    return box


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75,
         shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    if radius is not None:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    return sh


def pill(slide, x, y, w, h, label, fill=G_MAIN, color=WHITE, size=10.5):
    rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(slide, x + 0.14, y, w - 0.28, h, label, size=size, color=color, bold=True,
         anchor=MSO_ANCHOR.MIDDLE, spc=60)


def chevron(slide, x, y, w=0.20, h=0.34, color=G_MAIN):
    rect(slide, x, y, w, h, fill=color, shape=MSO_SHAPE.CHEVRON)


def rule(slide, x, y, w, h=0.014, color=LINE):
    rect(slide, x, y, w, h, fill=color)


def picture(slide, name, x, y, w, h, frame=False):
    pic = slide.shapes.add_picture(str(ASSETS / name), Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    if frame:                       # white figure card on the F6F8F9 page
        rect(slide, x, y, w, h, fill=None, line=LINE)
    return pic


def highlight(run, hexcolor="FFFFFF"):
    """White marker behind a run, as on the template's title slide."""
    rpr = run._r.get_or_add_rPr()
    hl = parse_xml(f'<a:highlight {nsdecls("a")}><a:srgbClr val="{hexcolor}"/></a:highlight>')
    fill = rpr.find(qn("a:solidFill"))
    if fill is not None:
        fill.addnext(hl)
    else:
        rpr.insert(0, hl)


def strip_placeholders(slide):
    for sh in list(slide.shapes):
        if sh.is_placeholder:
            sh._element.getparent().remove(sh._element)


def drop_all_slides(prs):
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)


# ------------------------------------------------------------------ pages
def band_page(prs, layout, title, tag, foot=True):
    s = prs.slides.add_slide(layout)
    strip_placeholders(s)
    text(s, 0.97, 0.13, 8.35, 0.60, title, size=22, bold=True, color=WHITE,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.45, 0.13, 3.63, 0.60, tag, size=9, color=BAND_TXT,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, spc=40)
    if foot:
        footer(s)
    return s


def footer(s):
    rule(s, 0.62, 7.06, 11.98)
    text(s, 0.62, 7.13, 8.6, 0.26,
         "Simulated IMU / GNSS corpus · 48 routes · 8.64 M rows · 8 held-out test routes",
         size=8, color=MUTE)
    text(s, 9.15, 7.13, 3.55, 0.26, "Urban navigation correction · Guo Haisu",
         size=8, color=MUTE, align=PP_ALIGN.RIGHT)


def quote(s, y, txt, h=0.70, size=15.5):
    rect(s, 0.83, y, 0.035, h, fill=G_MAIN)
    text(s, 1.04, y, 11.55, h, txt, size=size, font=F_SERIF, color=INK,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.04)


def stat(s, x, y, value, caption, w=3.55, val_size=26):
    text(s, x, y, w, 0.52, value, size=val_size, bold=True, color=G_DARK)
    text(s, x, y + 0.54, w, 0.62, caption, size=8.6, color=MUTE, spacing=1.05)


# ---------------------------------------------------------------- page 3.4
def page_34(prs, lay):
    s = band_page(prs, lay, "3.4  Training and Safety Gating",
                  "CHAPTER 03 · FOUNDATION MODEL AND LoRA")
    quote(s, 1.02, "One design decision and one safety mechanism: weight the loss "
                   "where the filter is blind, then let a one-standard-error gate "
                   "decide whether the correction ships.")

    # left: the loss design and the gate decision
    pill(s, 0.62, 1.88, 2.28, 0.40, "01 · LOSS DESIGN")
    text(s, 0.62, 2.36, 1.75, 0.90, "×22", size=54, bold=True, color=G_DARK,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.18, 2.50, 2.54, 0.70,
         "windows ×2, outage steps ×11 — 82 % of the gradient lands on the 7 % "
         "of the timeline where the filter is blind.", size=9.3, color=MUTE, spacing=1.1)

    for i, (v, lab) in enumerate([("7.0 %", "of the samples"),
                                  ("17.1 %", "of the squared residual"),
                                  ("81.7 %", "of the gradient")]):
        y = 3.42 + i * 0.60
        text(s, 0.62, y, 1.20, 0.42, v, size=20, bold=True, color=G_TEAL)
        text(s, 1.84, y + 0.04, 2.88, 0.42, lab, size=9.3, color=MUTE)

    rect(s, 0.62, 5.20, 4.10, 1.68, fill=G_DARK)
    text(s, 0.86, 5.34, 3.62, 0.26, "SAFETY FALLBACK · α = 0 ON 8 / 8 ROUTES",
         size=8.6, bold=True, color=G_PALE, spc=50)
    text(s, 0.86, 5.66, 3.62, 0.78,
         "Unguarded α* = 2.00 / 1.90 buys +2.7 % / +2.9 % on validation and loses "
         "2.2 % / 4.9 % on test. A one-standard-error rule returns α = 0.00 "
         "everywhere.", size=9.4, color=WHITE, spacing=1.12)
    rule(s, 0.86, 6.50, 3.62, 0.010, color=G_MAIN)
    text(s, 0.86, 6.58, 3.62, 0.24,
         "The gate is fitted per route on validation; test never sees it.",
         size=8.4, color=G_PALE)

    # right: the two evidence charts
    pill(s, 5.02, 1.88, 2.75, 0.40, "02 · TRAINING OUTCOME")
    text(s, 7.90, 1.94, 4.82, 0.30,
         "train loss −19 % / −66 % · validation stays above predict-zero",
         size=8.8, color=MUTE, align=PP_ALIGN.RIGHT)
    picture(s, "b_training_wide.png", 5.02, 2.32, 7.70, 2.07)

    pill(s, 5.02, 4.50, 2.90, 0.40, "03 · UNGUARDED CORRECTION")
    text(s, 8.02, 4.56, 4.70, 0.30, "α* chosen on validation only",
         size=8.8, color=MUTE, align=PP_ALIGN.RIGHT)
    picture(s, "c_gate_wide.png", 5.02, 4.94, 7.70, 1.93)
    return s


# ---------------------------------------------------------------- page 3.5
def page_35(prs, lay):
    s = band_page(prs, lay, "3.5  What It Bought — and Where It Stops",
                  "CHAPTER 03 · FOUNDATION MODEL AND LoRA")
    quote(s, 1.02, "Every tier removes what the one before cannot. What is left "
                   "is not noise — it is information the outage removed.")

    pill(s, 0.62, 1.88, 2.05, 0.40, "01 · THE LADDER")
    picture(s, "ladder.png", 0.55, 2.34, 6.25, 3.57)
    text(s, 0.62, 5.98, 6.15, 0.95,
         "B1 → B3 removes 90.5 % of the outage error; the worst single outage "
         "falls 413 → 28.3 m. The same smoother on GPU windows costs +8.2 % "
         "(10.75 m) — the price of seeing ±90 s instead of the whole route.",
         size=9.2, color=MUTE, spacing=1.12)

    pill(s, 7.10, 1.88, 1.92, 0.40, "02 · THE WALL")
    rows = [
        ("84 %", "of the RTS residual is a smooth quadratic in outage time — "
                 "a frozen initial-state error, not noise."),
        ("91 %", "of the outage error would vanish if that smooth part were known "
                 "exactly — that is, if the missing fix existed."),
        ("0.9 %", "is all an oracle IMU calibration adds; the bias is not what is "
                  "left over."),
        ("0.0 %", "is what the trained corrector recovers on unseen routes — the "
                  "channels do not carry the missing state."),
    ]
    for i, (v, txt) in enumerate(rows):
        y = 2.40 + i * 0.86
        text(s, 7.10, y, 1.28, 0.52, v, size=25, bold=True, color=G_DARK)
        text(s, 8.42, y + 0.02, 4.30, 0.78, txt, size=9.2, color=INK, spacing=1.1)

    rect(s, 7.10, 5.95, 5.62, 0.86, fill=G_DARK)
    text(s, 7.34, 6.06, 5.14, 0.26, "THE GATE IS THE RESULT  ·  α = 0 ON 8 / 8",
         size=9.0, bold=True, color=G_PALE, spc=50)
    text(s, 7.34, 6.32, 5.14, 0.42,
         "Direction transfers across routes (r ≈ 0.70); amplitude does not.",
         size=9.2, color=WHITE)
    return s


# ---------------------------------------------------------- chapter divider
def page_divider(prs, lay):
    s = prs.slides.add_slide(lay)
    strip_placeholders(s)
    # two lines, two sizes: the number in the sans face, the words in the serif
    text(s, 0.50, 2.83, 2.80, 1.90,
         [[("04", {"size": 60, "bold": True})],
          [("PART FOUR", {"size": 26, "font": F_SERIF, "spc": 600})]],
         color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
         spacing=1.15)

    text(s, 4.94, 2.06, 7.55, 1.10, "Spatial Information", size=48, bold=True,
         color=INK, anchor=MSO_ANCHOR.MIDDLE)

    sub = text(s, 5.66, 3.20, 7.0, 0.46, "", size=20, font=F_SERIF, color=INK)
    p = sub.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "What the Map Adds to a Time-Only Solution"
    _style_run(r, 20, INK, False, F_SERIF, False, None)
    highlight(r)

    bullets = ["Road-centreline projection", "Lane-level mapped result",
               "Three numbers, three limits"]
    for i, b in enumerate(bullets):
        y = 3.94 + i * 0.60
        tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(5.72),
                                 Inches(y + 0.06), Inches(0.20), Inches(0.17))
        tri.rotation = 90
        tri.shadow.inherit = False
        tri.fill.solid()
        tri.fill.fore_color.rgb = G_MAIN
        tri.line.fill.background()
        text(s, 6.12, y, 6.4, 0.46, b, size=20, color=INK)
    return s


# ---------------------------------------------------------------- page 4.1
def page_map(prs, lay):
    s = band_page(prs, lay, "4.1  The Road-Centreline Prior",
                  "CHAPTER 04 · SPATIAL INFORMATION")
    quote(s, 1.02, "Physics and learning both work in time. A map works in space — "
                   "it supplies exactly the component no inertial sensor can observe.")

    stat(s, 0.62, 1.86, "9.93 → 6.30 m",
         "test-mean outage error, B3 → map prior (−36.5 %)", w=3.55)
    stat(s, 4.40, 1.86, "8 / 8 routes", "improve, +26 % to +46 %", w=3.55)
    stat(s, 8.18, 1.86, "±1 m survey",
         "tolerated; at 2 m route 46 jumps to a parallel branch", w=3.55)
    for x in (4.28, 8.06):
        rect(s, x, 1.90, 0.012, 1.10, fill=LINE)

    picture(s, "fig7_map_prior_b3.png", 0.62, 2.92, 12.10, 3.30, frame=True)
    text(s, 0.62, 6.30, 12.10, 0.62,
         "Continuity-constrained projection — arc length monotone, speed limit, "
         "heading term — not nearest-point, which reaches only 7.52 m and regresses "
         "the self-intersecting route 46 by 36 %. Shown on the numpy reference (B3); "
         "the GPU trajectory follows the same path, 10.75 → 6.98 m.",
         size=8.8, color=MUTE, spacing=1.1)
    return s


# ---------------------------------------------------------------- page 4.2
def page_lane(prs, lay):
    s = band_page(prs, lay, "4.2  What It Means for the Map",
                  "CHAPTER 04 · SPATIAL INFORMATION")
    quote(s, 1.02, "A lane marking is observed relative to the vehicle, so a pose "
                   "error becomes a map error — the doubled, smeared lane line of a "
                   "map built from uncorrected trajectories.")

    stat(s, 0.62, 1.86, "43.9 → 6.5 → 0.28 m",
         "lane-line lateral offset: dead reckoning → filter → map prior", w=3.9)
    stat(s, 4.72, 1.86, "86.2 %",
         "of mapped points inside the 0.5 m HD tolerance", w=3.4)
    stat(s, 8.60, 1.86, "40.7 → 6.9 m",
         "along-track error, left untouched by the prior by design", w=3.9)
    for x in (4.60, 8.48):
        rect(s, x, 1.90, 0.012, 1.10, fill=LINE)

    picture(s, "fig8_laneline.png", 0.62, 2.92, 12.10, 3.38, frame=True)
    text(s, 0.62, 6.38, 12.10, 0.50,
         "0.28 m is the chord error of a 1 Hz centreline, not a pose accuracy: with "
         "a prior map this measures map update, not mapping from scratch.",
         size=8.8, color=MUTE, spacing=1.1)
    return s


# ---------------------------------------------------------------- page 4.3
def page_three(prs, lay):
    s = band_page(prs, lay, "4.3  Three Numbers, Three Limits",
                  "CHAPTER 04 · SPATIAL INFORMATION")
    quote(s, 1.02, "Each number belongs to a different limit — compute, geometry, "
                   "and mapping.")

    cards = [
        ("48.6×", "THROUGHPUT",
         "1 M IMU rows: 2,245 s on CPU vs 46.2 s on one RTX 4090; 177.6× at chunk "
         "500. The full 8.64 M-row corpus filters in 17 min."),
        ("90.5 %", "ACCURACY",
         "Outage error 104.65 → 9.93 m; worst outage 413 → 28.3 m. The windowed "
         "GPU pipeline pays 8.2 % for its parallelism."),
        ("0.28 m", "MAPPING",
         "Lane-line lateral offset after the prior; 86 % of mapped points inside "
         "the 0.5 m HD tolerance."),
    ]
    for i, (val, lab, body) in enumerate(cards):
        x = 0.62 + i * 3.89
        rect(s, x, 1.86, 3.66, 2.34, fill=WHITE, line=LINE,
             shape=MSO_SHAPE.FOLDED_CORNER)
        text(s, x + 0.28, 2.06, 3.10, 0.68, val, size=38, bold=True, color=G_DARK)
        text(s, x + 0.28, 2.82, 3.10, 0.26, lab, size=8.8, bold=True,
             color=G_TEAL, spc=80)
        text(s, x + 0.28, 3.14, 3.10, 0.92, body, size=9.6, color=INK, spacing=1.12)
        if i < 2:
            chevron(s, x + 3.70, 2.85)

    # what we would test next
    text(s, 0.62, 4.56, 5.0, 0.30, "WHAT WE WOULD TEST NEXT", size=9.2, bold=True,
         color=G_TEAL, spc=80)
    rule(s, 0.62, 4.90, 12.10)
    nxt = [
        "Replay a real IMU / GNSS log — the corpus here is simulated; the filter "
        "and the prior do not change, only the noise model does.",
        "Widen the window — 90 s of burn-in is a requirement, not the window "
        "length; a longer window recovers most of the 8.2 %.",
        "Feed the corrector the missing fix — until a sensor supplies what the "
        "outage removed, the gate keeps α at zero.",
    ]
    for i, t in enumerate(nxt):
        x = 0.62 + i * 3.89
        tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(x),
                                 Inches(5.12), Inches(0.13), Inches(0.11))
        tri.rotation = 90
        tri.shadow.inherit = False
        tri.fill.solid()
        tri.fill.fore_color.rgb = G_MAIN
        tri.line.fill.background()
        text(s, x + 0.24, 5.04, 3.36, 1.10, t, size=9.4, color=INK, spacing=1.14)

    rect(s, 0.62, 6.40, 12.10, 0.46, fill=WASH)
    text(s, 0.86, 6.40, 11.6, 0.46,
         [[("REPRODUCIBLE", {"bold": True, "color": G_TEAL, "spc": 60}),
           ("    48 routes · 8.64 M rows · every figure re-renders from results/summary.md",
            {"color": MUTE})]],
         size=8.6, anchor=MSO_ANCHOR.MIDDLE)
    return s


# ---------------------------------------------------------------- page thanks
def page_thanks(prs, lay):
    s = band_page(prs, lay, "Questions & Discussion",
                  "MS022161 · ENGINEERING BIG-DATA ANALYSIS", foot=False)
    text(s, 1.10, 2.00, 11.0, 1.0, "Thank You", size=54, bold=True, color=G_DARK)
    rule(s, 1.16, 3.24, 3.10, 0.03, color=G_MAIN)
    text(s, 1.16, 3.48, 10.9, 0.70,
         "GPU-Parallel Batched Kalman Smoothing and a LoRA-Tuned Time-Series "
         "Foundation Model", size=17, font=F_SERIF, color=INK, spacing=1.08)
    text(s, 1.16, 4.42, 10.8, 0.34,
         "Guo Haisu · MS022161 Engineering Big-Data Analysis · Southeast University",
         size=10.5, color=MUTE)
    return s


# ------------------------------------------------------------------- main
def main() -> None:
    # the deck is often still open in PowerPoint while it is being reviewed;
    # fall back to a -v2 sibling instead of dying on the locked file
    out = OUT
    try:
        shutil.copy2(TEMPLATE, out)
    except PermissionError:
        out = OUT.with_name(OUT.stem + "-v2" + OUT.suffix)
        shutil.copy2(TEMPLATE, out)
        print(f"{OUT.name} is locked -> writing {out.name}")
    prs = Presentation(out)
    master = prs.slide_masters[1]                    # the green master
    lay_band = master.slide_layouts[2]               # "节标题": band + seal + tab
    lay_div = master.slide_layouts[1]                # "标题和内容": left panel

    drop_all_slides(prs)
    page_34(prs, lay_band)
    page_35(prs, lay_band)
    page_divider(prs, lay_div)
    page_map(prs, lay_band)
    page_lane(prs, lay_band)
    page_three(prs, lay_band)
    page_thanks(prs, lay_band)

    # the same spoken script as deck pages 17-23, one page to one page
    notes_txt = pathlib.Path(__file__).resolve().parents[1] / "results" / "ppt" / "speaker_notes.txt"
    if notes_txt.exists():
        chunks = notes_txt.read_text(encoding="utf-8").split("--- slide ")[1:]
        notes = {int(c.split(" ---\n", 1)[0]): c.split(" ---\n", 1)[1].rstrip("\n")
                 for c in chunks}
        for page, src in enumerate(range(17, 24), 1):
            prs.slides[page - 1].notes_slide.notes_text_frame.text = notes[src]
        print("speaker notes copied: 7 pages")
    prs.save(out)
    print(f"{out.name}: {len(prs.slides._sldIdLst)} slides")

    dump = pathlib.Path(__file__).resolve().parents[1] / "results" / "ppt" / "green_deck_text.txt"
    with open(dump, "w", encoding="utf-8") as fh:
        for i, slide in enumerate(prs.slides, 1):
            fh.write(f"\n===== page {i} ({slide.slide_layout.name}) =====\n")
            for sh in slide.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    fh.write(sh.text_frame.text.replace("\n", " / ") + "\n")
    print(f"text dump -> {dump}")


if __name__ == "__main__":
    main()
